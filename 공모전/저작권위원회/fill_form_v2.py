"""공식 AI 창작기술서(docx) 2부 작성: 이미지 부문 / 숏폼 부문. 성명·서명·작성일은 비워 둠(출품자 직접 기재)."""
import os
import docx

os.chdir(os.path.dirname(os.path.abspath(__file__)))

# 기술서 수치는 v2 렌더 기록(render_params.json)과 최종 파일에서 직접 읽음(문서와 실제가 어긋나지 않도록)
import json, subprocess
RP = json.load(open("out/v2/render_params.json", encoding="utf-8"))
A = RP["audio"]
VID_DUR = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                "out/v2/숏폼_v2.mp4"], capture_output=True, text=True).stdout)
VOL_DB, TRIM, XF, END_LEN = A["volume_db"], A["song_trim_s"], A["crossfade_s"], A["ending_card_length_s"]
LOOP = A["ending_audio_source_s"]; LOOP_FADE = A["loop_end_fade_s"]
FONTS = "Bagel Fat One·Do Hyeon·Jua·Black Han Sans·Pretendard(모두 SIL OFL 1.1)"

def set_par(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]: r.text = ""
    else:
        p.add_run(text)

def check(doc, old, new):
    hit = 0
    pars = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in pars:
        if old in p.text:
            set_par(p, p.text.replace(old, new)); hit += 1
    assert hit, old

def fill_box(doc, ti, text):
    cell = doc.tables[ti].rows[0].cells[0]
    ps = cell.paragraphs
    for i, line in enumerate(text.split("\n")):
        if i < len(ps): set_par(ps[i], line)
        else: cell.add_paragraph(line)

TOOLS_IMG = "기타: Higgsfield 웹 – Nano Banana 2(이미지 생성), Claude(기획·대사 초안, 이미지 프롬프트 작성, 편집·조판 스크립트 작성)"
TOOLS_SF = ("기타: Suno AI v6(작곡·편곡·보컬 생성), Higgsfield 웹 – Nano Banana 2(이미지 생성), Claude(기획·가사 초안, 음악·이미지 프롬프트 작성, "
            "편집 스크립트 작성), faster-whisper·Demucs(가사 인식 검수·자막 타이밍 측정)")
PROMPT_IMG = """※ 아래는 실제 입력 프롬프트의 발췌입니다(… 표시는 생략, 전체 원문은 출품자가 보관한 프롬프트 기록에 있음).
[캐릭터 시트] "Original cartoon character design sheet for a Korean webtoon, two characters on a plain white background, each shown in front view and three-quarter view, full body, consistent design. Character 1: a small chubby creature whose body is two overlapping rounded rectangles like a soft pastel mint-green pair of sticky notes, big round shiny eyes… Character 2: a tiny fairy referee girl with short wavy lavender hair, black-and-white striped referee shirt, small translucent sparkly wings, a silver whistle on a lanyard, holding a red card… no text, no letters, no logos, no watermark."
[컷별 장면 – 캐릭터 시트를 참조 이미지로 사용] "Use the two characters from the reference image exactly as designed (…). Korean webtoon comic panel: (장면) … Plain unbranded laptop with no logo, plain walls with no posters or pictures of any characters … Absolutely no text, no letters, no words, no numbers, no logos, no brand marks, no watermark."
3컷 장면: "…points clearly up at a framed colorful abstract swirl painting hanging on the wall above the desk, as if telling the AI to copy that painting exactly; the laptop screen shows a small cute round light-blue robot assistant … looking worried…"
6컷 장면: "role-reversal finale, wide shot: the mint-green sticky-note creature wearing a little red beret stands proudly next to its own framed hand-drawn picture … a swarm of small floating speech-bubble comments with pink hearts and copy-arrow icons (no readable text) is flying toward the picture from a laptop screen … the fairy referee girl … raising a green card…"
※ 특정 작가·작품·캐릭터·브랜드 이름은 프롬프트에 사용하지 않았습니다."""
PROMPT_SF = """※ 실제 입력값(가사는 [Verse] 등 구조 태그만 빼고 전문을 그대로 적음).
[Suno 스타일] playful Korean novelty pop song, 122 BPM, bouncy synth bass, handclaps, retro chiptune internet sound effects, bright female lead vocal with kids group chant, catchy call-and-response chorus, comedic spoken intro, clear crisp Korean diction, clean modern mix, short under 45 seconds
[Suno 가사] Claude와 함께 작성하고 출품자가 확정한 가사(기존 노래 가사 미사용): "퍼가요~♡ ...잠깐만! / 마음에 쏙 든 사진 한 장, 노래 한 소절 / 퍼 가기 전에 잠깐 멈춰, 누가 만든 건지 봐요 / 출처만 쓰면 괜찮다? 아니죠, 허락이 먼저 / 무료로 쓰고 싶다면 공유마당 이용 조건 확인! / 퍼가요 말고 허락부터! (허락부터!) / 쓸 때는 출처까지! (출처까지!) / AI한테 시킬 때도 똑같이, 비슷하게는 NO! / 창작자에게 박수 짝짝, 케이 저작권 매너! / 오늘부터 우리 모두 저작권 매너인!"
[이미지] 6컷 만화와 같은 캐릭터 시트·컷별 장면 프롬프트 사용(이전 버전 6컷 원화 프롬프트 포함, 특정 작가·작품·캐릭터·브랜드 이름 미사용)."""

COMMON_CHECKS = ["□ 타인의 저작권", "□ 특정 작가", "□ 생성형 AI 결과물을 그대로", "□ AI 창작 기술서의 내용을", "□ 네, 최종 작품에"]

def make(kind):
    d = docx.Document("refs/AI_창작기술서.docx")
    if kind == "img":
        set_par(d.tables[1].rows[0].cells[1].paragraphs[0], "출처만 쓰면 괜찮다며?! (퍼가요 말고 허락부터!)")
        check(d, "□ 이미지", "■ 이미지")
        intro = ("'출처만 밝히면 남의 사진을 써도 괜찮다'는 흔한 오해를 바로잡고, 생성형 AI 시대의 올바른 저작물 이용 방법을 알리는 6컷 만화입니다.\n"
                 "남의 사진을 퍼 오며 '출처: 인터넷'만 적은 '복붙이'가 저작권 심판 요정에게 레드카드를 받고, AI에게 남의 그림을 똑같이 그려 달라고 시키다 다시 제지당한 뒤, "
                 "공유마당에서 이용 조건을 확인하고 직접 그림을 그리지만, 이번엔 자기 그림에 '퍼가요~♡' 댓글이 쏟아지자 \"허락부터 받아 줘요!\"라고 외치며 창작자의 입장을 깨닫는 역지사지 반전으로 끝납니다.\n"
                 "'출처 표시는 기본, 이용 허락·조건이 먼저'라는 핵심 메시지를 웃음 포인트와 함께 전달합니다.\n"
                 "※ 인용 등 저작권법상 예외와, CCL·공공누리처럼 표시된 조건만 지키면 이용할 수 있는 경우가 있다는 점을 고려해 '허락·조건'으로 표현했습니다. '똑같이·비슷하게'는 특정 작품의 창작적 표현을 닮게 유도하는 경우를 뜻하며, 화풍·아이디어가 비슷한 것까지 금지한다는 뜻은 아닙니다. 'AI 활용 사실 밝히기'는 개인 창작자에게 법적 의무가 아닌 투명성 매너로 제시했습니다(AI 기본법의 생성물 표시 의무는 AI 사업자 대상). 6컷의 '퍼 가기'는 무단 복제·게시를 뜻하며, 단순 링크 공유는 별개입니다.")
        tools, prompt = TOOLS_IMG, PROMPT_IMG
        use = ("캐릭터 시트를 먼저 생성해 두 캐릭터의 디자인을 고정하고, 이를 참조 이미지로 넣어 컷별 장면을 생성했습니다(6컷 반전 장면은 시안 2개 중 캐릭터 일관성이 맞는 것을 선택).\n"
               "생성 결과를 한 장씩 검수해 문제가 있는 컷은 프롬프트를 고쳐 다시 생성했습니다: 실제 브랜드 로고를 연상시키는 노트북 마크, 기존 캐릭터로 오인될 수 있는 벽 포스터, "
               "의도하지 않은 영문 글자, 손이 3개로 그려진 오류, 캐릭터 몸에 화면이 합쳐진 구도 오류.\n"
               "검수를 통과한 6컷을 골라 2단 3행 만화로 구성했습니다.")
        edit = ("이미지 생성 AI가 만든 글자는 쓰지 않고, 제목·대사·효과음·댓글 칩('퍼가요~♡')·이름표('그림: 복붙이')·사진 라벨('출처: 인터넷')·컷 번호를 편집 단계에서 별도로 조판했습니다(대사 초안은 Claude와 작성하고 출품자가 확정).\n"
                f"용도별 글꼴 체계를 설계해 적용하고({FONTS}), 화자별로 말풍선 형태를 나눴습니다(요정 대사 띠, 복붙이 타원 말풍선, 놀람 가시 풍선, 효과음 폭발 풍선).\n"
                "2컷의 영문 효과음은 한글 '삐익!!'으로 덮었고, 6컷 노트북 덮개의 작은 마크는 주변 색으로 지웠으며, AI 원화의 테두리는 잘라 내고 한 줄 테두리로 통일했습니다. 하단에 'AI 창작 매너 3' 띠와 필수 표기 문구를 넣었습니다.")
        contrib = ("이야기 구조(오해 → 제지 → AI 오용 → 제지 → 올바른 방법 → 역지사지 반전)와 캐릭터 설정(복사·붙여넣기를 상징하는 포스트잇 캐릭터 '복붙이', 레드카드·그린카드를 든 저작권 심판 요정)을 Claude와 함께 기획했고, "
                   "출품자가 제시된 안 중 콘셉트와 개선안(6컷 역지사지 반전 등)을 선택·승인했습니다.\n"
                   "생성 결과를 검수해 권리침해 소지가 있는 요소를 걸러 재생성·수정했고, 대사·효과음·글꼴·레이아웃을 구성해 완성했습니다.\n"
                   "기획·대사 초안, 이미지 프롬프트, 글꼴·말풍선·레이아웃 설계안, 생성 결과 검수, 마크 지우기·조판 등 편집 작업과 스크립트 작성은 Claude(AI)가 제안·수행했으며, 출품자는 콘셉트와 개선안 선택·승인, 대사 확정, 최종 결과 검토·확정을 맡았습니다.")
    else:
        set_par(d.tables[1].rows[0].cells[1].paragraphs[0], "퍼가요 말고 허락부터 (K-저작권 매너송)")
        check(d, "□ 숏폼 영상", "■ 숏폼 영상")
        intro = (f"옛 인터넷 유행어 '퍼가요~♡'를 훅으로 쓴 약 {round(VID_DUR)}초 저작권 매너송 숏폼입니다.\n"
                 "'출처만 쓰면 괜찮다?'는 오해를 바로잡아 이용 허락이 먼저라는 원칙을 전하고, 공유마당에서 이용 조건을 확인하는 방법, "
                 "AI에게 남의 작품을 똑같이·비슷하게 따라 하게 시키지 않는 AI 시대의 매너를 따라 부르기 쉬운 노래로 담았습니다.\n"
                 "※ 가사 '출처만 쓰면 괜찮다? 아니죠, 허락이 먼저'는 일반 원칙을 쉽게 전한 표현이며, 인용 등 법정 예외와 CCL·공공누리처럼 표시된 조건을 지키면 이용할 수 있는 경우도 있습니다. '똑같이, 비슷하게는 NO!'는 특정 작품의 창작적 표현을 닮게 유도하는 경우를 뜻하며 화풍·아이디어 유사는 해당하지 않고, 'AI 활용 사실 밝히기'는 개인 창작자에게 법적 의무가 아닌 투명성 매너로 제시했습니다(AI 기본법의 생성물 표시 의무는 AI 사업자 대상).")
        tools, prompt = TOOLS_SF, PROMPT_SF
        use = ("확정한 가사로 Suno에서 두 버전을 생성하고, 보컬을 분리해 음성 인식으로 가사 전달력을 비교했습니다. 이후 발음이 불분명한 구절('복붙 전에')과 표현을 보완한 가사('퍼 가기 전에', '똑같이, 비슷하게는 NO!')로 다시 두 버전을 생성해, 핵심 가사가 정확히 들리는 버전을 골랐습니다.\n"
               "화면에는 6컷 만화용으로 생성한 장면(반전 장면 이전 버전의 6컷 원화 포함)을 썼고, 가사 내용에 맞는 장면을 골라 배치했습니다.")
        edit = (f"음원은 음량을 {VOL_DB}dB 올리고 곡 끝 잔향 일부를 잘라({TRIM:.2f}초까지 사용) {XF*1000:.0f}ms 크로스페이드로 엔딩과 이은 것 외에는 생성된 곡 그대로 사용했습니다. 엔딩(약 {END_LEN:.1f}초)에는 같은 곡의 후렴 구간({LOOP[0]:.2f}~{LOOP[1]:.2f}초, '퍼가요 말고 허락부터!')을 한 번 더 썼고, 반복 재생을 위해 끝 {LOOP_FADE*1000:.0f}ms만 페이드했습니다(새 음원·효과음 없음).\n"
                f"가사 자막을 입력해 노래 시작 시점에 맞춰 배치하고, 박자에 맞춘 컷 전환·확대·플래시·흔들림 효과와 'AI 창작 매너 3' 진행 표시, 엔딩 카드를 설계했습니다(글꼴: {FONTS}).\n"
                "영문 효과음을 한글로 바꾸고 배경의 작은 마크를 지웠으며, '퍼가요~♡ 출처: 인터넷 ㅎ' 댓글 말풍선·'그림: 복붙이' 이름표·'짝 짝!' 효과음 글자를 추가했습니다. 필수 표기 문구는 그림 바로 아래 띠에 처음부터 끝까지 상시 넣었습니다.")
        contrib = ("'퍼가요~♡' 유행어를 반전시킨 노래 콘셉트와 가사 구성(오해 바로잡기 → 허락·출처 → 공유마당 → AI 매너)을 Claude와 함께 기획했고, 출품자가 콘셉트·가사·개선안을 확정했습니다.\n"
                   "음원 후보를 가사 전달력 기준으로 선택하고, 장면 배치·자막·연출을 구성해 숏폼을 완성했습니다.\n"
                   "기획·가사 초안, 음악·이미지 프롬프트, 글꼴·레이아웃·모션 설계안, 가사 인식 검수와 자막 싱크 측정, 편집 작업과 스크립트 작성은 Claude(AI)가 제안·수행했으며, 출품자는 콘셉트와 개선안 선택·승인, 가사 확정, 음원 선택 승인, 최종 결과 검토·확정을 맡았습니다.")
    check(d, "□ 올바른 저작물 이용", "■ 올바른 저작물 이용")
    check(d, "□ K-저작권 매너", "■ K-저작권 매너")
    check(d, "□ 기타: ________________________________________________________________", "■ " + tools)
    for c in COMMON_CHECKS: check(d, c, c.replace("□", "■"))
    import re
    for p in d.tables[5].rows[1].cells[0].paragraphs:                               # '기타' 줄 끝 밑줄 제거
        if "■ 기타" in p.text: set_par(p, re.sub(r"_+\s*$", "", p.text))
    fill_box(d, 4, intro); fill_box(d, 7, prompt); fill_box(d, 9, use); fill_box(d, 11, edit); fill_box(d, 13, contrib)
    os.makedirs("out/v2", exist_ok=True)
    out = f"out/v2/AI창작기술서_{'이미지부문_출처만쓰면괜찮다며' if kind == 'img' else '숏폼부문_퍼가요말고허락부터'}.docx"
    d.save(out); print("saved", out)

make("img"); make("sf")
