"""공식 AI 창작기술서(docx) 2부 작성: 이미지 부문 / 숏폼 부문. 성명·서명·작성일은 비워 둠(출품자 직접 기재)."""
import os
import docx

os.chdir(os.path.dirname(os.path.abspath(__file__)))

# 기술서 수치는 최종 렌더·스크립트에서 직접 읽음(문서와 실제가 어긋나지 않도록)
import re, subprocess
_src = open("make_short.py", encoding="utf-8").read()
VID_DUR = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                "최종제출/숏폼부문_퍼가요말고허락부터.mp4"], capture_output=True, text=True).stdout)
VOL_DB = re.search(r"volume=([0-9.]+)dB", _src).group(1)
FADE_D = re.search(r"afade=t=out:st=\{DUR - [0-9.]+:\.2f\}:d=([0-9.]+)", _src).group(1)
TAIL = re.search(r"TOTAL = DUR \+ ([0-9.]+)", _src).group(1)
TOP_NOTE = re.search(r"if t < ([0-9.]+):\s+# 초반 상단 표기", _src).group(1).rstrip("0").rstrip(".")

def set_par(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]: r.text = ""
    else:
        p.add_run(text)

def check(doc, old, new):
    hit = 0
    pars = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in pars:
        if old in p.text:
            set_par(p, p.text.replace(old, new)); hit += 1
    assert hit, old

def fill_box(doc, ti, text):
    cell = doc.tables[ti].rows[0].cells[0]
    ps = cell.paragraphs
    for i, line in enumerate(text.split("\n")):
        if i < len(ps): set_par(ps[i], line)
        else: cell.add_paragraph(line)

TOOLS_IMG = "기타: Higgsfield 웹 – Nano Banana 2(이미지 생성), Claude(기획·대사 초안, 이미지 프롬프트 작성, 편집 스크립트 작성)"
TOOLS_SF = ("기타: Suno AI v6(작곡·편곡·보컬 생성), Higgsfield 웹 – Nano Banana 2(이미지 생성), Claude(기획·가사 초안, 음악·이미지 프롬프트 작성, "
            "편집 스크립트 작성), faster-whisper·Demucs(가사 인식 검수·자막 타이밍 측정)")
PROMPT_IMG = """[캐릭터 시트] Original cartoon character design sheet for a Korean webtoon, two characters on a plain white background, front and three-quarter view. Character 1: a small chubby creature whose body is two overlapping rounded rectangles like a pastel mint-green pair of sticky notes… Character 2: a tiny fairy referee girl with short wavy lavender hair, black-and-white striped referee shirt, small translucent wings, a silver whistle, holding a red card… Bright flat 2D webtoon illustration, no text, no logos.
[컷별 장면 – 캐릭터 시트를 참조 이미지로 사용] "Use the two characters from the reference image exactly as designed. Korean webtoon comic panel: (장면 설명)… Plain unbranded laptop with no logo, plain walls with no posters or pictures of any characters, no text, no logos."
장면 설명 예) 3컷: the sticky-note creature points at a framed abstract painting made by someone else as if asking the robot on the laptop screen to copy it exactly; the robot looks worried.
※ 특정 작가·작품·캐릭터·브랜드 이름은 프롬프트에 사용하지 않았습니다."""
PROMPT_SF = """[Suno 스타일] playful Korean novelty pop song, 122 BPM, bouncy synth bass, handclaps, retro chiptune internet sound effects, bright female lead vocal with kids group chant, catchy call-and-response chorus, comedic spoken intro, short under 45 seconds
[Suno 가사] Claude와 함께 작성하고 출품자가 확정한 가사 전문을 입력(기존 노래 가사 미사용): "퍼가요~♡ …잠깐만! / 마음에 쏙 든 사진 한 장, 노래 한 소절 / 복붙 전에 잠깐 멈춰, 누가 만든 건지 봐요 / 출처만 쓰면 괜찮다? 아니죠, 허락이 먼저 / … / 퍼가요 말고 허락부터! / 쓸 때는 출처까지! / AI한테 시킬 때도 남의 작품 그대로는 NO! / 창작자에게 박수 짝짝, K-저작권 매너!"
[이미지] 6컷 만화와 같은 캐릭터 시트·컷별 장면 프롬프트 사용(특정 작가·작품·캐릭터·브랜드 이름 미사용)."""

COMMON_CHECKS = ["□ 타인의 저작권", "□ 특정 작가", "□ 생성형 AI 결과물을 그대로", "□ AI 창작 기술서의 내용을", "□ 네, 최종 작품에"]

def make(kind):
    d = docx.Document("refs/AI_창작기술서.docx")
    if kind == "img":
        set_par(d.tables[1].rows[0].cells[1].paragraphs[0], "출처만 쓰면 괜찮다며?! (복붙이의 K-저작권 매너 대작전)")
        check(d, "□ 이미지", "■ 이미지")
        intro = ("'출처만 밝히면 남의 사진을 써도 괜찮다'는 흔한 오해를 바로잡고, 생성형 AI 시대의 올바른 저작물 이용 방법을 알리는 6컷 만화입니다.\n"
                 "남의 사진을 퍼 오며 '출처: 인터넷'만 적은 '복붙이'가 저작권 심판 요정에게 레드카드를 받고, AI에게 남의 그림을 똑같이 그려 달라고 시키다 다시 제지당한 뒤, "
                 "공유마당에서 이용 조건을 확인하고 직접 창작하며 AI 사용 사실까지 표기하는 '창작이'로 성장하는 이야기입니다.\n"
                 "'출처 표시는 기본, 이용 허락·조건이 먼저'라는 핵심 메시지를 웃음 포인트와 함께 전달합니다.\n"
                 "※ 인용 등 저작권법상 예외와, CCL·공공누리처럼 표시된 조건만 지키면 이용할 수 있는 경우가 있다는 점을 고려해 '허락·조건'으로 표현했습니다.")
        tools, prompt = TOOLS_IMG, PROMPT_IMG
        use = ("캐릭터 시트를 먼저 생성해 두 캐릭터의 디자인을 고정하고, 이를 참조 이미지로 넣어 컷별 장면 6장을 생성했습니다.\n"
               "생성 결과를 한 장씩 검수해 문제가 있는 컷은 프롬프트를 고쳐 다시 생성했습니다: 실제 브랜드 로고를 연상시키는 노트북 마크, 기존 캐릭터로 오인될 수 있는 벽 포스터, "
               "의도하지 않은 영문 글자, 손이 3개로 그려진 오류, 캐릭터 몸에 화면이 합쳐진 구도 오류.\n"
               "검수를 통과한 6컷을 골라 2단 3행 만화로 구성했습니다.")
        edit = ("이미지 생성 AI가 만든 글자는 쓰지 않고, 제목·대사·효과음·컷 번호를 편집 단계에서 별도로 입력했습니다(대사 초안은 Claude와 작성하고 출품자가 확정).\n"
                "2컷의 영문 효과음은 한글 '삐익!!'으로 덮어 바꿨고, 6컷 배경 노트북 덮개의 작은 마크는 주변 색으로 지웠습니다.\n"
                "3컷은 화자가 분명하도록 대사를 말풍선 하나로 정리했고, 하단에 필수 표기 문구와 사용 폰트(Noto Sans CJK KR, SIL OFL)를 넣었습니다.")
        contrib = ("이야기 구조(오해 → 제지 → AI 오용 → 제지 → 올바른 방법 → 성장)와 캐릭터 설정(복사·붙여넣기를 상징하는 포스트잇 캐릭터 '복붙이', 레드카드·그린카드를 든 저작권 심판 요정)을 Claude와 함께 기획했고, "
                   "출품자가 제시된 안 중 콘셉트를 선택·승인했습니다.\n"
                   "생성 결과를 검수해 권리침해 소지가 있는 요소를 걸러 재생성·수정했고, 대사·효과음·레이아웃을 구성해 완성했습니다.\n"
                   "기획·대사 초안, 프롬프트와 편집 스크립트 작성에는 Claude(AI)를 활용했으며, 출품자는 콘셉트 선택, 대사 확정, 최종 결과 검토·확정을 맡았습니다.")
    else:
        set_par(d.tables[1].rows[0].cells[1].paragraphs[0], "퍼가요 말고 허락부터 (K-저작권 매너송)")
        check(d, "□ 숏폼 영상", "■ 숏폼 영상")
        intro = (f"옛 인터넷 유행어 '퍼가요~♡'를 훅으로 쓴 약 {round(VID_DUR)}초 저작권 매너송 숏폼입니다.\n"
                 "'출처만 쓰면 괜찮다?'는 오해를 바로잡아 이용 허락이 먼저라는 원칙을 전하고, 공유마당에서 이용 조건을 확인하는 방법, "
                 "AI에게 남의 작품을 그대로 따라 하게 시키지 않는 AI 시대의 매너를 따라 부르기 쉬운 노래로 담았습니다.\n"
                 "※ 가사 '출처만 쓰면 괜찮다? 아니죠, 허락이 먼저'는 일반 원칙을 쉽게 전한 표현이며, 인용 등 법정 예외와 CCL·공공누리처럼 표시된 조건을 지키면 이용할 수 있는 경우도 있습니다.")
        tools, prompt = TOOLS_SF, PROMPT_SF
        use = ("확정한 가사로 Suno에서 두 버전을 생성하고, 보컬을 분리해 음성 인식으로 가사 전달력을 비교했습니다. 그중 핵심 가사('허락부터', 'NO!', '공유마당')가 정확히 들리는 버전을 골랐습니다.\n"
               "화면에는 6컷 만화용으로 생성한 장면과 캐릭터 시트를 썼고, 가사 내용에 맞는 장면을 골라 배치했습니다.")
        edit = (f"음원은 음량을 {VOL_DB}dB 올리고 곡 끝 {FADE_D}초 페이드아웃과 엔딩 무음(약 {TAIL}초)을 더한 것 외에는 생성된 곡 그대로 사용했습니다.\n"
                "가사 자막을 입력해 노래 시작 시점에 맞춰 배치하고, 장면별 확대·이동 효과와 전환을 설계했습니다.\n"
                f"영문 효과음을 한글로 바꾸고 배경의 작은 마크를 지웠으며, 1컷 화면 속 빈 라벨에 '출처: 인터넷'을 써 넣고 제목·엔딩 문구를 추가했습니다. 필수 표기 문구는 그림 바로 아래에 상시 넣고 초반 {TOP_NOTE}초에는 상단에도 넣었습니다.")
        contrib = ("'퍼가요~♡' 유행어를 반전시킨 노래 콘셉트와 가사 구성(오해 바로잡기 → 허락·출처 → 공유마당 → AI 매너)을 Claude와 함께 기획했고, 출품자가 콘셉트와 가사를 확정했습니다.\n"
                   "음원 후보를 가사 전달력 기준으로 선택하고, 장면 배치·자막·연출을 구성해 숏폼을 완성했습니다.\n"
                   "기획·가사 초안, 프롬프트와 편집 스크립트 작성에는 Claude(AI)를 활용했으며, 출품자는 콘셉트 선택, 가사 확정, 최종 결과 검토·확정을 맡았습니다.")
    check(d, "□ 올바른 저작물 이용", "■ 올바른 저작물 이용")
    check(d, "□ K-저작권 매너", "■ K-저작권 매너")
    check(d, "□ 기타: ________________________________________________________________", "■ " + tools)
    for c in COMMON_CHECKS: check(d, c, c.replace("□", "■"))
    import re
    for p in d.tables[5].rows[1].cells[0].paragraphs:                               # '기타' 줄 끝 밑줄 제거
        if "■ 기타" in p.text: set_par(p, re.sub(r"_+\s*$", "", p.text))
    fill_box(d, 4, intro); fill_box(d, 7, prompt); fill_box(d, 9, use); fill_box(d, 11, edit); fill_box(d, 13, contrib)
    os.makedirs("최종제출", exist_ok=True)
    out = f"최종제출/AI창작기술서_{'이미지부문_출처만쓰면괜찮다며' if kind == 'img' else '숏폼부문_퍼가요말고허락부터'}.docx"
    d.save(out); print("saved", out)

make("img"); make("sf")
